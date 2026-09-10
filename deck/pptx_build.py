# -*- coding: utf-8 -*-
"""Builds an editable PowerPoint from the deck.

Text is native and editable. The vector pieces (the map, the funnel, the
product screens) come in as pictures, because they were never text to begin
with. Run: python3 pptx_build.py
"""
import copy, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

PX = 9525                      # one CSS pixel in EMU at 96 dpi
W, H = 1600, 900
ART = 'pptx-art'

G  = RGBColor(0x0B, 0x0A, 0x09)   # ground
T  = RGBColor(0xF7, 0xF3, 0xEC)   # type
M  = RGBColor(0x7C, 0x76, 0x6D)   # muted
A  = RGBColor(0xFF, 0x5C, 0x1A)   # accent
L  = RGBColor(0x24, 0x1F, 0x1B)   # line
P  = RGBColor(0x14, 0x12, 0x10)   # panel
RED = RGBColor(0xE8, 0x50, 0x3F)

KUFI, BODY, MONO = 'Thmanyah', 'Thmanyah', 'IBM Plex Mono'
NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    return prs


def blank(prs, num=None, cap=None, foot=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, G)
    if num:
        txt(s, 104, 84, 400, num, 15, M, font=MONO, align='r')
        txt(s, 1096, 84, 400, cap or '', 18, M, align='l')
    if foot and num:
        rect(s, 104, 800, 1392, 1, L)
        txt(s, 104, 812, 400, num + ' / 18', 14, L, font=MONO, align='r')
        txt(s, 1096, 810, 400, 'كامن', 16, M, bold=True, align='l')
    return s


def rect(s, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x*PX), Emu(y*PX), Emu(w*PX), Emu(h*PX))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.text_frame.text = ''
    return sh


def txt(s, x, y, w, text, size, color=T, bold=False, font=BODY, align='r',
        h=None, line=1.35, rtl=True):
    """align: r = flush right (RTL natural), l = flush left, c = centred."""
    box = s.shapes.add_textbox(Emu(x*PX), Emu(y*PX), Emu(w*PX), Emu((h or size*2)*PX))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {'r': PP_ALIGN.RIGHT, 'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER}[align]
    pPr = p._p.get_or_add_pPr()
    if rtl:
        pPr.set('rtl', '1')
    ln = etree.SubElement(pPr, NS + 'lnSpc')
    etree.SubElement(ln, NS + 'spcPct').set('val', str(int(line * 100000)))
    r = p.add_run(); r.text = text
    f = r.font
    f.size = Pt(size * 0.75)          # css px to points
    f.bold = bold
    f.color.rgb = color
    f.name = font
    # Arabic needs the complex-script typeface set too, or it falls back
    rPr = r._r.get_or_add_rPr()
    for tag in ('cs', 'ea'):
        el = etree.SubElement(rPr, NS + tag)
        el.set('typeface', font)
    return box


def pic(s, name, x, y, w, h=None):
    path = name if os.path.exists(name) else os.path.join(ART, name)
    kw = {'width': Emu(w*PX)}
    if h:
        kw['height'] = Emu(h*PX)
    return s.shapes.add_picture(path, Emu(x*PX), Emu(y*PX), **kw)


def TX(s, y, text, size, color=T, bold=False, font=BODY, h=None, line=1.35,
       x=104, w=1392, align='r'):
    """Full-width RTL text, flush right like the deck."""
    return txt(s, x, y, w, text, size, color, bold, font, align, h, line)


def build():
    prs = deck()

    # 01 cover
    s = blank(prs, foot=False)
    rect(s, 1156, 88, 340, 12, A)
    TX(s, 190, 'كامن', 224, T, True, KUFI, h=270, line=1.16)
    TX(s, 470, 'الموجود الغير مكتشف', 26, M)
    TX(s, 508, 'K A M I N', 24, A, font=MONO)
    TX(s, 566, 'خريطة حيّة للقدرات المحلية.', 52, T, True, KUFI, h=80)
    rect(s, 104, 776, 1392, 1, L)
    txt(s, 104, 790, 400, '01 / 18', 17, L, font=MONO, align='r')
    TX(s, 788, 'هاكاثون ابتكار · المسار الأول', 22, M, w=600, x=896, align='l')

    # 02 the question
    s = blank(prs, '02', 'الافتتاحية')
    rect(s, 1346, 168, 150, 10, A)
    TX(s, 194, '«هل فيه مورد محلي؟»', 96, T, True, KUFI, h=130)
    TX(s, 330, 'سؤال تظن إجابته سهلة، لدرجة إنك تتوقّع تلقاها بواحدة من هذي:', 25, M)
    pic(s, 'cards.png', 104, 386, 1392)
    TX(s, 704, 'وفي الحقيقة؟ محد يقدر يجاوبه.', 50, A, True, KUFI, h=72)

    # 03 the three reasons
    s = blank(prs, '03', 'لماذا يهم')
    TX(s, 250, 'شركات الصندوق تشتري من موردين خارجيين، لثلاثة أسباب.', 56, T, True, KUFI, h=170, line=1.3)
    for i, (tag, head, body, hot) in enumerate([
            ('السبب الأول', 'لسبب وجيه', 'فعلاً محد يصنعه محلياً.', False),
            ('السبب الثاني', 'لأنه ما يناسب', 'بسبب جودته، وقت التسليم، أو اختلاف المعايير.', True),
            ('السبب الثالث', 'لسبب مؤسف', 'فيه أحد يصنعه على بُعد ساعتين، بس محد يدري عنه.', True)]):
        x = 1030 - i * 494
        rect(s, x, 448, 466, 240, RGBColor(0x1A, 0x0E, 0x09) if hot else P, A if hot else L)
        txt(s, x + 30, 480, 406, tag, 16, A if hot else M)
        txt(s, x + 30, 512, 406, head, 42, A if hot else T, bold=True, font=KUFI, h=60)
        txt(s, x + 30, 578, 406, body, 26, T if hot else M, h=100, line=1.6)

    # 04 the problem, three faces
    s = blank(prs, '04', 'المشكلة')
    TX(s, 160, 'ليه؟', 66, T, True, KUFI, h=90)
    TX(s, 250, 'ليه ما نعرفه، وليه غير مناسب.', 24, M)
    pic(s, 'problem.png', 104, 306, 1392)

    # 05 the interstitial
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, A)
    D = RGBColor(0x5A, 0x22, 0x06)
    txt(s, 104, 84, 400, '05', 17, D, font=MONO, align='r')
    txt(s, 1096, 84, 400, 'الفصل الثاني', 18, D, align='l')
    txt(s, 200, 320, 1200, 'والحل؟', 200, G, bold=True, font=KUFI, align='c', h=270, line=1.2)
    rect(s, 104, 800, 1392, 1, RGBColor(0x8A, 0x3A, 0x10))
    txt(s, 104, 812, 400, '05 / 18', 16, D, font=MONO, align='r')
    txt(s, 1096, 810, 400, 'كامن', 18, D, bold=True, align='l')

    # 06 what KAMIN is
    s = blank(prs, '06', 'وش هو كامن؟')
    pic(s, '../KAMIN.png', 1364, 232, 132, 132)
    txt(s, 944, 240, 400, 'كامن', 88, T, bold=True, font=KUFI, align='r', h=110, line=1.05)
    txt(s, 944, 350, 400, 'الموجود الغير مكتشف', 22, M, align='r')
    txt(s, 944, 384, 400, 'K A M I N', 17, A, font=MONO, align='r', rtl=False)
    rect(s, 916, 232, 1, 170, L)
    txt(s, 104, 236, 790, 'منصّة وكلاء ذكاء اصطناعي تشتغل لحالها، على مدار الساعة.', 40, T,
        bold=True, font=KUFI, align='r', h=120, line=1.35)
    txt(s, 104, 360, 790, 'تبني للصندوق خريطة حيّة للقدرات المحلية، وتربطها بطلب شركات المحفظة مجمّعاً، '
        'وتكمل الطريق للمورّد نفسه.', 22, M, align='r', h=90, line=1.65)
    for i, head in enumerate(['مؤتمتة', 'متعددة الوكلاء', 'ما توقف']):
        x = 1054 - i * 476
        rect(s, x, 508, 442, 2, A)
        txt(s, x, 526, 442, head, 30, T, bold=True, font=KUFI, h=44)
        if i == 1:
            for j, f in enumerate(['fig-coordinator.png', 'fig-detective.png',
                                   'fig-auditor.png', 'fig-advisor.png']):
                pic(s, f, x + 350 - j * 96, 578, 84)
    rect(s, 104, 700, 1392, 1, L)
    txt(s, 1396, 736, 100, 'والمُخرج', 21, M, align='r')
    for i, t in enumerate(['موردون موثّقون', 'سجل الفجوات', 'نسبة التغطية']):
        x = 942 - i * 424
        rect(s, x, 720, 404, 62, RGBColor(0x1E, 0x10, 0x0A), A)
        txt(s, x, 736, 404, t, 26, A, bold=True, font=KUFI, align='c')

    # 07 the team, and the turn
    s = blank(prs, '07', 'طيب كيف نحلها؟')
    TX(s, 168, 'ببساطة. نوظّف فريق.', 58, T, True, KUFI, h=84)
    for i, (img, name) in enumerate([('fig-coordinator.png', 'المنسّق'), ('fig-detective.png', 'المحقّق'),
                                     ('fig-auditor.png', 'المدقّق'), ('fig-advisor.png', 'المستشار')]):
        x = 1163 - i * 353
        rect(s, x, 268, 333, 3, A)
        pic(s, img, x + 96, 286, 141)
        txt(s, x, 484, 333, name, 32, T, bold=True, font=KUFI, align='c', h=46)
    rect(s, 104, 566, 1392, 1, L)
    txt(s, 1000, 592, 496, 'احنا ما نوظّفهم.', 46, M, bold=True, font=KUFI, align='r', h=66)
    txt(s, 640, 578, 340, 'نشغّلهم.', 84, A, bold=True, font=KUFI, align='r', h=112, line=1.1)
    xr = 1496
    for chip, hot, w in [('وكلاء ذكاء اصطناعي', True, 216), ('دقيقتين للمصنع بدل يوم', False, 250),
                         ('٢٤ ساعة بلا توقف', False, 196), ('ما ينسون ولا معلومة', False, 214)]:
        xr -= w
        rect(s, xr, 706, w, 46, RGBColor(0x2A, 0x14, 0x0B) if hot else None, A if hot else L)
        txt(s, xr, 719, w, chip, 19, A if hot else M, align='c')
        xr -= 12

    # 08 the wall, and why now
    s = blank(prs, '08', 'فلماذا محد سواها؟')
    txt(s, 900, 246, 596, 'وليه محد سواها للآن؟', 36, T, bold=True, font=KUFI, align='r', h=52)
    for i, (a_, b_) in enumerate([('مصانع المملكة، ٢٠٢٥', '١٢٩٤٦ مصنع'), ('وكل مصنع يحتاج', 'يوم كامل')]):
        y = 316 + i * 66
        rect(s, 900, y, 596, 1, L)
        txt(s, 1196, y + 18, 300, a_, 21, M, align='r')
        txt(s, 900, y + 14, 300, b_, 30, T, bold=True, font=KUFI, align='l')
    rect(s, 900, 450, 596, 1, A)
    txt(s, 1370, 484, 126, 'يعني', 21, M, align='r')
    txt(s, 900, 464, 450, '٣٥ سنة', 76, A, bold=True, font=KUFI, align='l', h=104, line=1.1)
    txt(s, 900, 596, 596, 'وأول ما تخلص، بتكون معلوماتك آوت ديتد!', 30, T, bold=True, font=KUFI,
        align='r', h=90, line=1.4)
    rect(s, 866, 246, 1, 400, L)
    txt(s, 150, 246, 680, 'وليش الحين؟', 36, A, bold=True, font=KUFI, align='r', h=52)
    txt(s, 150, 300, 680, 'لأن اللي كان يحتاج باحث ليوم كامل، صار وكيل ذكاء اصطناعي يسويه:', 22, M,
        align='r', h=64, line=1.6)
    for i, t in enumerate(['يقرأ الموقع والسجل والترسيات والكتالوجات', 'دقيقتين للمصنع بدل يوم كامل',
                           '٢٤ ساعة بلا توقف، وبالتوازي', 'وما ينسى ولا معلومة لقاها']):
        txt(s, 150, 376 + i * 38, 680, t + '  ·', 21, T, align='r')
    rect(s, 150, 556, 680, 1, A)
    txt(s, 150, 578, 680, 'و٣٥ سنة تصير أقل من عشر ساعات.', 34, A, bold=True, font=KUFI, align='r', h=52)

    # 09 the map
    s = blank(prs, '09', 'ما الذي يبنونه')
    TX(s, 228, 'خريطة كامن للقدرات المحلية', 52, T, True, KUFI, h=74)
    pic(s, 'map.png', 104, 320, 1392)

    # 10 zoom, discovery
    s = blank(prs, '10', 'زووم · الاكتشاف')
    TX(s, 118, 'وش يصير قبل ما يوصل الخريطة؟', 44, T, True, KUFI, h=62)
    pic(s, 'minimap.png', 104, 182, 1392)
    for i, (head, sub, lines, out, img) in enumerate([
            ('مصادر عامة', '', ['منافسات · ترميز · صنع في السعودية  ·', 'التجارة · الموارد البشرية · المقاولين  ·'],
             'اسم ورقم سجل  →', None),
            ('المحقّق', 'يبحث، ويتواصل', ['يقرأ الموقع والسجل والترسيات  ·', 'ويتواصل مع اللي يلقاه  ·'],
             'مورّد ومعه أدلته  →', 'bust-detective.png'),
            ('المدقّق', 'يتحقّق', ['يدقّق في معلومات المحقّق  ·', 'يصنّف ويوثّق المورّد  ·'],
             'قدرة موثّقة تدخل الخريطة  →', 'bust-auditor.png')]):
        x = 1046 - i * 471
        rect(s, x, 276, 450, 226, None if i == 0 else P, L)
        if img:
            pic(s, img, x + 16, 292, 50)
        txt(s, x + 20, 294, 410, head, 30, T, bold=True, font=KUFI, align='r')
        if sub:
            txt(s, x + 20, 334, 410, sub, 16, A, align='r')
        for j, ln in enumerate(lines):
            txt(s, x + 20, 372 + j * 34, 410, ln, 20, T, align='r')
        txt(s, x + 20, 456, 410, out, 20, A, align='r')
    txt(s, 1276, 524, 220, 'تواصل المحقّق', 28, T, bold=True, font=KUFI, align='r')
    txt(s, 700, 530, 556, 'وكل رسالة فيها رابط التسجيل في مساهمة.', 18, M, align='r')
    pic(s, 'funnel.png', 104, 566, 1392)

    # 11 zoom, the decision
    s = blank(prs, '11', 'زووم · القرار')
    TX(s, 118, 'ولما الخريطة تقول محد يقدر؟', 44, T, True, KUFI, h=62)
    pic(s, 'minimap2.png', 104, 182, 1392)
    x = 1496
    for head, sub, w, dashed in [('الخريطة', '', 300, True),
                                 ('المستشار', 'ما يوقف عند «محد يقدر». يجرّب ست درجات بالترتيب.', 470, False),
                                 ('سجل الفجوات', 'اللي ما انحل، مقيّد بقيمته السنوية', 274, False),
                                 ('نسبة التغطية', 'اللي قدرنا نغطّيه محلياً، بعد كل درجة', 274, False)]:
        x -= w
        rect(s, x, 270, w, 128, None if dashed else RGBColor(0x1A, 0x0E, 0x09), L if dashed else A)
        txt(s, x + 18, 292, w - 36, head, 27, T if dashed else A, bold=True, font=KUFI, align='r')
        if sub:
            txt(s, x + 18, 330, w - 36, sub, 17, M, align='r', h=62, line=1.4)
        x -= 22
    for i, (n, name, tool, hot) in enumerate([('١', 'نشتريه', 'أمر شراء', False), ('٢', 'نقسّمه', 'ترسية مقسّمة', False),
                                              ('٣', 'نستثمر', 'توسعة مقابل تعاقد', True), ('٤', 'نشارك', 'شراكة ونقل معرفة', True),
                                              ('٥', 'نوطّن', 'استثمار مباشر', True), ('٦', 'نستورد', 'استيراد', False)]):
        col, row = i % 3, i // 3
        x = 1042 - col * 469
        y = 428 + row * 118
        rect(s, x, y, 454, 102, RGBColor(0x1A, 0x0E, 0x09) if hot else P, A if hot else L)
        txt(s, x + 396, y + 34, 44, n, 17, A if hot else M, align='c')
        txt(s, x + 180, y + 26, 210, name, 34, A if hot else T, bold=True, font=KUFI, align='r')
        rect(s, x + 22, y + 30, 152, 42, None, A if hot else L)
        txt(s, x + 22, y + 42, 152, tool, 18, A if hot else T, align='c')

    # 12 the map again
    s = blank(prs, '12', 'ما الذي يبنونه')
    TX(s, 228, 'خريطة كامن للقدرات المحلية', 52, T, True, KUFI, h=74)
    pic(s, 'map.png', 104, 320, 1392)

    # 13 behind the scenes
    s = blank(prs, '13', 'خلف الكواليس')
    TX(s, 168, 'وش يشغّل الخريطة؟', 40, T, True, KUFI, h=56)
    txt(s, 380, 178, 600, 'نفس الخريطة، والمنقّط تحتها هو اللي يغذّي الوكلاء', 18, M, align='l')
    txt(s, 104, 178, 260, '●  وكيل ذكاء اصطناعي', 17, M, align='r')
    pic(s, 'engine.png', 104, 226, 1392)

    # 14 the supplier list
    s = blank(prs, '14', 'المُخرج')
    pic(s, 'suppliers.png', 104, 147, 1392)

    # 15 the math
    s = blank(prs, '15', 'الحساب')
    txt(s, 1096, 244, 400, 'الطريقة القديمة', 17, M, align='l')
    txt(s, 596, 244, 400, 'كامن', 17, A, align='l')
    for i, (k, a_, b_) in enumerate([('الوقت', '٣٥ سنة', 'أقل من ١٠ ساعات'),
                                     ('التكلفة', 'ملايين الريالات', 'أقل من ٤ آلاف دولار'),
                                     ('الحداثة', 'تقرير يطلع قديم', 'يحدّث نفسه')]):
        y = 286 + i * 90
        rect(s, 104, y, 1392, 1, L)
        txt(s, 1296, y + 26, 200, k, 29, M, bold=True, font=KUFI, align='r')
        txt(s, 896, y + 28, 380, a_, 27, M, align='l')
        txt(s, 496, y + 28, 380, b_, 27, T, bold=True, align='l')
    TX(s, 596, 'نفس الشغل. بس بساعات، مو بسنوات.', 56, T, True, KUFI, h=84)

    # 16 KAMIN today: what it built, and what it cost
    s = blank(prs, '16', 'كامن اليوم')
    TX(s, 160, 'كامن اليوم.', 62, T, True, KUFI, h=88)
    TX(s, 252, 'الأرقام من قاعدة البيانات، والتكلفة من الفواتير.', 24, M)
    rect(s, 796, 316, 1, 452, L)
    txt(s, 1096, 318, 400, 'وش بناه', 17, A, align='l')
    for i, (n, lab) in enumerate([
            ('١٥,٠٢٥', 'مورّد على الخريطة، منهم ١٥٢ ما هم في ترميز'),
            ('٥١,٩٩٢', 'قدرة موثّقة: مورّد واحد، منتج واحد، بمواصفة'),
            ('٤,٨٣٦', 'رمز تعريفة مسجّل، من ١٠,٧٩٤ في التصنيف'),
            ('١,١٤٦', 'تشغيلة وكيل، في ٧٣ ساعة تشغيل')]):
        y = 352 + i * 96
        rect(s, 836, y, 660, 1, L)
        txt(s, 1306, y + 20, 190, n, 42, A if i == 0 else T, bold=True, align='r')
        txt(s, 836, y + 32, 450, lab, 21, M, align='l', h=60, line=1.5)

    txt(s, 396, 318, 360, 'وش كلّف', 17, A, align='l')
    rect(s, 104, 352, 660, 1, L)
    txt(s, 424, 366, 340, 'صفر ريال', 64, A, bold=True, font=KUFI, align='r', h=88)
    txt(s, 104, 396, 310, 'على ٢.٩٣ مليون توكن دخل', 21, M, align='l')
    txt(s, 104, 460, 660, 'qwen3.5:9b يشتغل محلياً على Ollama، وTavily على باقته المجانية '
        'ونتائجه مخزّنة، فإعادة التشغيل ما تكلف شي.', 21, M, h=76, line=1.7)
    txt(s, 104, 556, 660, 'ولو شغّلناها عند مزوّد ثاني، نفس التوكنات:', 17, M)
    for i, (name, price, hot) in enumerate([('Ollama · qwen3.5:9b', '$0', True),
                                            ('Gemini Flash', '$2', False),
                                            ('GPT-4o', '$12', False),
                                            ('Claude Opus', '$81', False)]):
        y = 592 + i * 46
        rect(s, 104, y, 660, 1, L)
        txt(s, 404, y + 12, 360, name, 17, T if hot else M, font=MONO, align='r', rtl=False)
        txt(s, 104, y + 11, 200, price, 18, A if hot else T, bold=True, font=MONO, align='l', rtl=False)

    # 17 the close
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, G)
    rect(s, 1156, 236, 340, 12, A)
    TX(s, 292, 'مساهمة تعرف مين سجّل.', 86, M, True, KUFI, h=120, line=1.3)
    TX(s, 404, 'كامن يعرف اللي ما رفع يده.', 86, T, True, KUFI, h=120, line=1.3)
    TX(s, 556, 'احنا مو بديل للمنصة. احنا الطبقة اللي تحتها.', 30, M)

    # 18 who we are
    s = blank(prs, '18', 'من نحن')
    rect(s, 1316, 158, 180, 10, A)
    TX(s, 186, 'ثلاثة أشخاص. وآلاف الوكلاء.', 76, T, True, KUFI, h=112, line=1.3)
    TX(s, 298, 'الفريق اللي يبني كامن.', 26, M)
    for i, (img, name, role) in enumerate([('face-mohammed.jpg', 'محمد شلبي', 'AI-Native Principle Engineer'),
                                           ('face-ali.jpg', 'علي باموالم', 'Product Consultant'),
                                           ('face-abdulaziz.jpg', 'عبدالعزيز الحارثي', 'Data Engineer / DevOps')]):
        x = 1232 - i * 460
        pic(s, img, x, 372, 264, 264)
        txt(s, x - 60, 662, 384, name, 38, T, bold=True, font=KUFI, align='c', h=54)
        txt(s, x - 60, 716, 384, role, 20, M, align='c', rtl=False)

    return prs


if __name__ == '__main__':
    out = '/Users/mohammedshalabi/workspace/personal/KAMIN-Deck-editable.pptx'
    build().save(out)
    print('wrote', out, os.path.getsize(out), 'bytes')
